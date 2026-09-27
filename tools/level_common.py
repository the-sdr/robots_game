"""Shared helpers for the level-map pipeline (level_design/level_map.xlsx).

Grid: 1 m cells, cell (x, z) is centred on world (x, z). North (-Z) is up the sheet.
Cell tokens: CODE[+Y][@HEADING], several per cell separated by spaces.
"""
import json
import math
import os
import re

import numpy as np
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET = os.path.join(ROOT, "level_design", "level_map.xlsx")
ASSETS_JSON = os.path.join(ROOT, "level_design", "build", "assets.json")

# code -> (asset name, kind). kind decides how the Godot baker instances it.
CODES = {
    "TC1": ("CommonTree_1", "tree"), "TC3": ("CommonTree_3", "tree"), "TC5": ("CommonTree_5", "tree"),
    "TD1": ("DeadTree_1", "tree"), "TT2": ("TwistedTree_2", "tree"), "TT4": ("TwistedTree_4", "tree"),
    "RM1": ("Rock_Medium_1", "rock"),
    "BCO": ("Bush_Common", "bush"), "BCF": ("Bush_Common_Flowers", "bush"),
    "BFE": ("Fern_1", "prop"), "BFL": ("Flower_3_Group", "prop"),
    "WUB": ("Wall_UnevenBrick_Straight", "wall"),
    "VN1": ("Prop_Vine1", "prop"), "VN2": ("Prop_Vine2", "prop"),
    "SPS": ("RockPath_Round_Wide", "stone"),
    "SC1": ("Building_Small_1", "building"), "SC2": ("Building_Medium_2_001", "building"),
    "SC3": ("Building_Large_2", "building"),
    "CP1": ("collectible", "collectible"), "CP2": ("collectible", "collectible"),
    "CP3": ("collectible", "collectible"), "CP4": ("collectible", "collectible"),
}
# Fixed in world.tscn or pure markers: never built from the sheet.
IGNORED = {"SHO", "SCH", "SPN", "X", "P"}
TREE_CODES = {"TC1": "CommonTree_1", "TC3": "CommonTree_3", "TC5": "CommonTree_5",
              "TD1": "DeadTree_1", "TT2": "TwistedTree_2", "TT4": "TwistedTree_4"}

# Must match scripts/solid_tree.gd TRUNKS (radius, local centre offset x, z).
TRUNKS = {
    "CommonTree_1": (0.55, 0.02, 0.11), "CommonTree_3": (0.55, 0.01, 0.11),
    "CommonTree_5": (0.50, 0.07, 0.10), "DeadTree_1": (0.52, 0.13, 0.0),
    "TwistedTree_2": (1.12, 0.06, -0.01), "TwistedTree_4": (1.15, 0.02, 0.10),
}
ROBOT_RADIUS = 0.38
WALL_SIZE = (2.0, 3.12, 0.41)          # measured Wall_UnevenBrick_Straight
WALL_CENTRE_Z = -0.11                  # local z of the wall's centre

# House (fixed in world.tscn): outer box of the hand-built collision walls.
HOUSE_MIN = (-2.79, -7.84)
HOUSE_MAX = (2.79, -2.21)
HOUSE_ROOF = 4.3


def load_assets():
    return json.load(open(ASSETS_JSON))


TOKEN_RE = re.compile(r"^([A-Z]+\d*)(?:\+(-?\d+(?:\.\d+)?))?(?:@(\d+))?$")


def parse_token(tok):
    m = TOKEN_RE.match(tok)
    if not m:
        raise ValueError("Unrecognised cell token %r" % tok)
    return m.group(1), float(m.group(2) or 0.0), int(m.group(3) or 0)


def format_token(code, yoff=0.0, heading=0):
    s = code
    if abs(yoff) >= 0.05:
        s += "+%g" % round(yoff, 1)
    if heading % 360:
        s += "@%d" % (heading % 360)
    return s


def heading_to_basis(heading_deg):
    """Row-major 3x3 basis whose -Z points at compass heading (0 = N = -Z, 90 = E = +X)."""
    t = -math.radians(heading_deg)
    c, s = math.cos(t), math.sin(t)
    return [c, 0.0, s, 0.0, 1.0, 0.0, -s, 0.0, c]


class Sheet:
    """Reads Layout + Height. Keeps the workbook so callers can write back."""

    def __init__(self, path=SHEET):
        self.path = path
        self.wb = load_workbook(path)
        self.layout = self.wb["Layout"]
        self.height = self.wb["Height"]
        ws = self.layout
        self.xs = [ws.cell(1, c).value for c in range(2, ws.max_column + 1)]
        self.zs = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        self.x_min, self.x_max = min(self.xs), max(self.xs)
        self.z_min, self.z_max = min(self.zs), max(self.zs)

    def cell(self, ws, x, z):
        return ws.cell(z - self.z_min + 2, x - self.x_min + 2)

    def tokens(self, x, z):
        v = self.cell(self.layout, x, z).value
        return str(v).split() if v not in (None, "") else []

    def set_tokens(self, x, z, toks):
        self.cell(self.layout, x, z).value = " ".join(toks) if toks else None

    def cells(self):
        for x in self.xs:
            for z in self.zs:
                toks = self.tokens(x, z)
                if toks:
                    yield x, z, toks

    def anchors(self):
        out = {}
        for x in self.xs:
            for z in self.zs:
                v = self.cell(self.height, x, z).value
                if v not in (None, ""):
                    out[(x, z)] = float(v)
        return out

    def save(self):
        self.wb.save(self.path)


def smoothstep(e0, e1, v):
    t = np.clip((v - e0) / (e1 - e0), 0.0, 1.0)
    return t * t * (3 - 2 * t)


def sheet_heights(sheet):
    """Height for every sheet cell: filled cells fixed, empty cells blended
    (harmonic interpolation = each empty cell is the average of its neighbours).
    Returns array [z_index, x_index] with z_index 0 = sheet.z_min."""
    anchors = sheet.anchors()
    nz, nx = len(sheet.zs), len(sheet.xs)
    h = np.zeros((nz, nx))
    fixed = np.zeros((nz, nx), dtype=bool)
    for (x, z), v in anchors.items():
        h[z - sheet.z_min, x - sheet.x_min] = v
        fixed[z - sheet.z_min, x - sheet.x_min] = True
    if not fixed.any():
        return h
    # initial guess: inverse-distance weighting, then relax
    az, ax = np.nonzero(fixed)
    av = h[fixed]
    zz, xx = np.mgrid[0:nz, 0:nx]
    d2 = (zz[..., None] - az) ** 2 + (xx[..., None] - ax) ** 2 + 1e-6
    w = 1.0 / d2 ** 2
    guess = (w * av).sum(-1) / w.sum(-1)
    h = np.where(fixed, h, guess)
    for _ in range(4000):
        p = np.pad(h, 1, mode="edge")
        avg = (p[:-2, 1:-1] + p[2:, 1:-1] + p[1:-1, :-2] + p[1:-1, 2:]) * 0.25
        h = np.where(fixed, h, avg)
    # Harmonic blending leaves a sharp point at each lone anchor; a light blur
    # over everything rounds those off (anchors end up within a few cm).
    for _ in range(SMOOTH_PASSES):
        p = np.pad(h, 1, mode="edge")
        h = (p[:-2, :-2] + p[:-2, 1:-1] + p[:-2, 2:] + p[1:-1, :-2] + p[1:-1, 1:-1] +
             p[1:-1, 2:] + p[2:, :-2] + p[2:, 1:-1] + p[2:, 2:]) / 9.0
    return h


SMOOTH_PASSES = 2


def terrain_grid(sheet, x0=-110, x1=110, z0=-150, z1=80):
    """Heights on a 1 m grid covering the sheet plus the background. Outside the
    sheet: the nearest sheet-edge height plus a rise, so the far forest climbs
    away and the horizon line is hidden."""
    inner = sheet_heights(sheet)
    xs = np.arange(x0, x1 + 1)
    zs = np.arange(z0, z1 + 1)
    X, Z = np.meshgrid(xs, zs)
    cx = np.clip(X, sheet.x_min, sheet.x_max)
    cz = np.clip(Z, sheet.z_min, sheet.z_max)
    base = inner[(cz - sheet.z_min).astype(int), (cx - sheet.x_min).astype(int)]
    d_out = np.hypot(X - cx, Z - cz)
    rolling = 0.8 * np.sin(0.13 * X + 0.4) * np.cos(0.11 * Z + 1.1) + 0.5 * np.sin(0.23 * X - 0.19 * Z)
    rise = 7.0 * smoothstep(0.0, 45.0, d_out) + rolling * smoothstep(0.0, 15.0, d_out)
    return {"x0": int(x0), "z0": int(z0), "nx": len(xs), "nz": len(zs),
            "heights": (base + rise).astype(float)}


def height_at(grid, x, z):
    """Bilinear sample of a terrain_grid."""
    h = grid["heights"]
    fx, fz = x - grid["x0"], z - grid["z0"]
    ix, iz = int(math.floor(fx)), int(math.floor(fz))
    ix = min(max(ix, 0), grid["nx"] - 2)
    iz = min(max(iz, 0), grid["nz"] - 2)
    tx, tz = fx - ix, fz - iz
    a = h[iz, ix] * (1 - tx) + h[iz, ix + 1] * tx
    b = h[iz + 1, ix] * (1 - tx) + h[iz + 1, ix + 1] * tx
    return float(a * (1 - tz) + b * tz)


def min_height_around(grid, x, z, radius):
    """Lowest ground within radius, so wide objects never float on a slope."""
    best = height_at(grid, x, z)
    if radius <= 0:
        return best
    for a in range(8):
        ang = a * math.pi / 4
        best = min(best, height_at(grid, x + math.cos(ang) * radius, z + math.sin(ang) * radius))
    return best


def point_segment_distance(px, pz, ax, az, bx, bz):
    dx, dz = bx - ax, bz - az
    L = dx * dx + dz * dz
    t = 0.0 if L == 0 else max(0.0, min(1.0, ((px - ax) * dx + (pz - az) * dz) / L))
    return math.hypot(px - (ax + t * dx), pz - (az + t * dz))


def box_distance(px, pz, lo, hi):
    dx = max(lo[0] - px, 0.0, px - hi[0])
    dz = max(lo[1] - pz, 0.0, pz - hi[1])
    return math.hypot(dx, dz)


def cell_hash(x, z, salt=0):
    """Deterministic pseudo-random int per cell (stable across runs)."""
    n = (x * 73856093) ^ (z * 19349663) ^ (salt * 83492791)
    n = (n ^ (n >> 13)) * 1274126177
    return (n ^ (n >> 16)) & 0x7FFFFFFF
